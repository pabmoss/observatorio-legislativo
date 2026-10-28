#!/usr/bin/env python3
"""Genera una ficha Word editable por proyecto desde data/proyectos.json."""
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
import json
ROOT=Path(__file__).resolve().parents[1]
DATA=json.loads((ROOT/'data'/'proyectos.json').read_text(encoding='utf-8'))
NAVY=RGBColor(22,47,59); MUTED=RGBColor(95,108,113)

def hyperlink(p,label,url):
  part=p.part;rid=part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
  link=OxmlElement('w:hyperlink'); link.set(qn('r:id'),rid); run=OxmlElement('w:r')
  pr=OxmlElement('w:rPr'); color=OxmlElement('w:color'); color.set(qn('w:val'),'27566B');pr.append(color)
  u=OxmlElement('w:u');u.set(qn('w:val'),'single');pr.append(u);run.append(pr)
  tt=OxmlElement('w:t');tt.text=label;run.append(tt);link.append(run);p._p.append(link)

def generate(r):
  d=Document();s=d.sections[0];s.page_height=Inches(11.7);s.page_width=Inches(8.3)
  s.top_margin=Inches(.54);s.bottom_margin=Inches(.56);s.left_margin=Inches(.67);s.right_margin=Inches(.67)
  styles=d.styles;norm=styles['Normal'];norm.font.name='Aptos';norm.font.size=Pt(8.7);norm.font.color.rgb=NAVY
  norm.paragraph_format.space_after=Pt(4)
  for nm,sz in [('FichaTitulo',13),('FichaSeccion',9.2)]:
    sty=styles.add_style(nm,WD_STYLE_TYPE.PARAGRAPH);sty.font.name='Aptos';sty.font.size=Pt(sz);sty.font.bold=True;sty.font.color.rgb=NAVY
    sty.paragraph_format.space_before=Pt(9 if nm=='FichaSeccion' else 2);sty.paragraph_format.space_after=Pt(4)
  p=d.add_paragraph();p.style='Normal';run=p.add_run('OBSERVATORIO LEGISLATIVO  /  DESPACHO DEL H.S. CARLOS EDUARDO ENRÍQUEZ CAICEDO')
  run.font.size=Pt(7);run.font.bold=True;run.font.color.rgb=MUTED
  d.add_paragraph(r['radicado'],style='FichaTitulo')
  p=d.add_paragraph(r['titulo']);p.style='Normal';p.runs[0].font.bold=True;p.runs[0].font.size=Pt(10)
  tab=d.add_table(rows=0,cols=2);tab.alignment=WD_TABLE_ALIGNMENT.CENTER
  tab.columns[0].width=Inches(1.22);tab.columns[1].width=Inches(5.5)
  for label,key in [('Tipo de ley','tipo_ley'),('Estado documentado','estado'),('Origen','origen'),('Comisión','comision'),('Sesión y actividad','sesion')]:
    cells=tab.add_row().cells
    for cell in cells: cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cells[0].text=label;cells[1].text=r.get(key,'')
    for run in cells[0].paragraphs[0].runs: run.font.bold=True;run.font.size=Pt(8);run.font.color.rgb=MUTED
    for run in cells[1].paragraphs[0].runs:run.font.size=Pt(8)
  for heading,key in [('Breve exposición de motivos','motivos'),('Objeto del texto radicado','objeto'),('Análisis del soporte actual','analisis'),('Ponentes','ponentes')]:
    if not r.get(key):continue
    d.add_paragraph(heading,style='FichaSeccion')
    for block in r[key].split('\n\n'):d.add_paragraph(block)
  d.add_paragraph('Fuentes y documentos oficiales',style='FichaSeccion')
  for link in r['documentos']:
    p=d.add_paragraph(style='Normal');p.paragraph_format.space_after=Pt(2)
    hyperlink(p,link['etiqueta'],link['url'])
  if r['notas_documentales']:
    d.add_paragraph('Notas de verificación',style='FichaSeccion')
    for text in r['notas_documentales']:d.add_paragraph(text)
  foot=s.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
  rr=foot.add_run('Corte documental: '+DATA['corte_documental']+'  ·  Una citación o ponencia no acredita aprobación')
  rr.font.size=Pt(7);rr.font.color.rgb=MUTED
  out=ROOT/r['documento_word'];out.parent.mkdir(exist_ok=True,parents=True);d.save(out)
  return out

if __name__=='__main__':
 for r in DATA['proyectos']:generate(r)
 print('Generados',len(DATA['proyectos']),'Word editables')
